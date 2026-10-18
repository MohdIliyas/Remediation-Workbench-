"""Write reviewed alt text (or "decorative") back into the output files."""
from __future__ import annotations

import os
from collections import defaultdict

from docx import Document
from lxml import etree

from .common import docx_to_tagged_pdf
from .pdf_figures import make_figures_artifacts, set_figure_alt

WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
ADEC_NS = "http://schemas.microsoft.com/office/drawing/2017/decorative"
DECORATIVE_URI = "{C183D7F6-B498-43B3-948B-1728B52AA6E4}"


def _set_decorative(docpr, on: bool):
    ext_lst = docpr.find(f"{{{A_NS}}}extLst")
    existing = None
    if ext_lst is not None:
        for ext in ext_lst.findall(f"{{{A_NS}}}ext"):
            if ext.get("uri") == DECORATIVE_URI:
                existing = ext
    if on and existing is None:
        if ext_lst is None:
            ext_lst = etree.SubElement(docpr, f"{{{A_NS}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{A_NS}}}ext", uri=DECORATIVE_URI)
        etree.SubElement(ext, f"{{{ADEC_NS}}}decorative", nsmap={"adec": ADEC_NS}, val="1")
    elif not on and existing is not None:
        ext_lst.remove(existing)
        if len(ext_lst) == 0:
            docpr.remove(ext_lst)


def _apply_docx(path: str, decisions: dict[int, dict]):
    doc = Document(path)
    for n, docpr in enumerate(doc.element.body.iter(f"{{{WP_NS}}}docPr"), 1):
        d = decisions.get(n)
        if not d:
            continue
        if d["status"] == "decorative":
            docpr.set("descr", "")
            _set_decorative(docpr, True)
        else:
            docpr.set("descr", d["text"])
            _set_decorative(docpr, False)
    doc.save(path)


def _apply_pptx(path: str, decisions: dict[tuple, dict]):
    from pptx import Presentation
    from .pptx_fix import _cnvpr, _iter_shapes
    prs = Presentation(path)
    for n, slide in enumerate(prs.slides, 1):
        for sh in _iter_shapes(slide.shapes):
            d = decisions.get((n, sh.shape_id))
            if not d:
                continue
            cnv = _cnvpr(sh)
            if cnv is None:
                continue
            if d["status"] == "decorative":
                cnv.set("descr", "")
                _set_decorative(cnv, True)
            else:
                cnv.set("descr", d["text"])
                _set_decorative(cnv, False)
    prs.save(path)


def apply_decisions(outdir: str, items: list[dict], decisions: list[dict], lang: str) -> list[str]:
    """Update items in place and rewrite the files they point to. Returns notes for the user."""
    by_id = {it["id"]: it for it in items}
    docx_jobs: dict[str, dict[int, dict]] = defaultdict(dict)
    pdf_jobs: dict[str, dict[int, str]] = defaultdict(dict)
    pdf_artifacts: dict[str, set] = defaultdict(set)
    pptx_jobs: dict[str, dict[tuple, dict]] = defaultdict(dict)
    regen: dict[str, str] = {}
    notes = []

    for d in decisions:
        it = by_id.get(d.get("id"))
        status = d.get("status")
        text = " ".join((d.get("text") or "").split())[:1000]
        if not it or status not in ("approved", "decorative"):
            continue
        if status == "approved" and not text:
            continue
        it["status"], it["final"] = status, (text if status == "approved" else "")
        tgt = it["target"]
        if tgt["kind"] == "docx":
            docx_jobs[tgt["file"]][tgt["index"]] = {"status": status, "text": text}
            if tgt.get("pdf"):
                regen[tgt["file"]] = tgt["pdf"]
        elif tgt["kind"] == "pptx":
            pptx_jobs[tgt["file"]][(tgt["slide"], tgt["shape_id"])] = {"status": status, "text": text}
            if tgt.get("pdf"):
                regen[tgt["file"]] = tgt["pdf"]
        elif tgt["kind"] == "pdf-figure":
            if status == "approved":
                pdf_jobs[tgt["file"]][tgt["index"]] = text
            else:
                pdf_artifacts[tgt["file"]].add(tgt["index"])

    for name, dec in docx_jobs.items():
        _apply_docx(os.path.join(outdir, name), dec)
    for name, dec in pptx_jobs.items():
        _apply_pptx(os.path.join(outdir, name), dec)
    for docx_name, pdf_name in regen.items():
        from .pdf_fix import fix_metadata  # local import avoids a cycle
        pdf_path = os.path.join(outdir, pdf_name)
        if docx_name.lower().endswith(".pptx"):
            from .pipeline import pptx_to_tagged_pdf
            pptx_to_tagged_pdf(os.path.join(outdir, docx_name), pdf_path)
        else:
            docx_to_tagged_pdf(os.path.join(outdir, docx_name), pdf_path)
        fix_metadata(pdf_path, "", lang)
        notes.append(f"Regenerated {pdf_name}")
    for name, dec in pdf_jobs.items():
        set_figure_alt(os.path.join(outdir, name), dec)
    # artifacts last: removing figures from the tree would renumber them for set_figure_alt
    for name, idxs in pdf_artifacts.items():
        n = make_figures_artifacts(os.path.join(outdir, name), idxs)
        notes.append(f"Marked {n} decorative figure(s) as artifacts in {name}")
    return notes
