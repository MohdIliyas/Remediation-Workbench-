"""Scanned PDF -> structured Word document, using Tesseract's layout output.

Tesseract reports each word's position and height. Words are grouped into
paragraphs, and paragraphs set noticeably larger than the body text become
headings. The result is a clean, reflowable text version that exports to a
tagged PDF. Figures and equations are not carried over (they are flagged).
"""
from __future__ import annotations

import csv
import glob
import os
import statistics
import subprocess
import tempfile
from collections import defaultdict

from docx import Document
from docx.shared import Pt

from .structure import build_structure, clean_leaders


def _ocr_page(png: str, lang: str) -> list[dict]:
    from .common import check_stop, run_cmd
    check_stop()
    r = run_cmd(["tesseract", png, "-", "-l", lang, "--psm", "3", "tsv"],
                       capture_output=True, text=True, timeout=600)
    rows = csv.DictReader(r.stdout.splitlines(), delimiter="\t", quoting=csv.QUOTE_NONE)
    words = []
    for row in rows:
        try:
            if row["level"] != "5" or not row["text"].strip() or float(row["conf"]) < 0:
                continue
            words.append({
                "key": (int(row["block_num"]), int(row["par_num"])),
                "line": int(row["line_num"]),
                "h": int(row["height"]),
                "left": int(row["left"]), "top": int(row["top"]), "w": int(row["width"]),
                "conf": float(row["conf"]),
                "text": row["text"].strip(),
            })
        except (KeyError, ValueError):
            continue
    return words


def _mean_conf(words: list[dict]) -> float:
    return sum(w["conf"] for w in words) / len(words) if words else 0.0


def _osd_rotation(png: str) -> tuple[int, float]:
    """Tesseract's guess at the clockwise turn that makes the page upright, and its confidence."""
    from .common import check_stop, run_cmd
    check_stop()
    r = run_cmd(["tesseract", png, "-", "--psm", "0"], capture_output=True, text=True, timeout=300)
    rot, conf = 0, 0.0
    for line in r.stdout.splitlines():
        if line.startswith("Rotate:"):
            rot = int(line.split(":")[1]) % 360
        elif line.startswith("Orientation confidence:"):
            conf = float(line.split(":")[1])
    return rot, conf


def ocr_review(png: str, turn: int, words: list[dict]) -> dict:
    """What a person should proofread on this page: low-confidence words, and printed areas the
    OCR found no text in (missed text, a table, or a picture that needs alt text)."""
    low = [w["text"] for w in words if w["conf"] < 60 and sum(ch.isalnum() for ch in w["text"]) >= 3]
    uncovered = 0
    try:
        import cv2
        import numpy as np
        from PIL import Image
        with Image.open(_rotated_copy(png, turn)) as im:
            g = np.asarray(im.convert("L"))
        h, w = g.shape
        ink = (g < 140).astype("uint8")
        for wd in words:  # blank out everything the OCR did read
            x0, y0 = max(wd["left"] - 6, 0), max(wd["top"] - 6, 0)
            ink[y0:wd["top"] + wd["h"] + 6, x0:wd["left"] + wd["w"] + 6] = 0
        # drop long thin rules and borders, then look for sizeable unread inked areas
        lines = cv2.bitwise_or(
            cv2.morphologyEx(ink, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(w // 8, 20), 1))),
            cv2.morphologyEx(ink, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(h // 8, 20)))))
        rest = cv2.dilate(np.clip(ink.astype("int16") - lines, 0, 1).astype("uint8"), np.ones((25, 25), "uint8"))
        n, _, stats, _ = cv2.connectedComponentsWithStats(rest, connectivity=8)
        for i in range(1, n):
            bw, bh = stats[i, cv2.CC_STAT_WIDTH], stats[i, cv2.CC_STAT_HEIGHT]
            if bw * bh > 0.015 * w * h and min(bw, bh) > 0.03 * min(w, h):
                uncovered += 1
    except Exception:
        pass
    return {"low_words": low[:8], "low_count": len(low), "uncovered": uncovered}


def _rotated_copy(png: str, cw: int) -> str:
    """Save the page image turned `cw` degrees clockwise (PIL turns counter-clockwise)."""
    if cw % 360 == 0:
        return png
    from PIL import Image
    out = png[:-4] + f"_r{cw}.png"
    with Image.open(png) as im:
        im.rotate(-cw, expand=True).save(out)
    return out


def ocr_upright(png: str, lang: str) -> tuple[list[dict], int, float]:
    """OCR a page the right way up. Returns (words, clockwise turn applied, mean confidence).

    Uses Tesseract's orientation detection; if the result still reads poorly, tries the
    other three orientations and keeps the best one."""
    guess, _ = _osd_rotation(png)
    best_words = _ocr_page(_rotated_copy(png, guess), lang)
    best_rot, best_conf = guess, _mean_conf(best_words)
    if best_conf < 70 or len(best_words) < 5:
        for cw in (0, 90, 180, 270):
            if cw == guess:
                continue
            words = _ocr_page(_rotated_copy(png, cw), lang)
            conf = _mean_conf(words)
            if conf > best_conf + 5 and len(words) >= 5:
                best_words, best_rot, best_conf = words, cw, conf
    return best_words, best_rot, best_conf


def _paragraphs(words: list[dict]) -> list[dict]:
    paras = defaultdict(lambda: defaultdict(list))
    order = []
    for w in words:
        if w["key"] not in paras:
            order.append(w["key"])
        paras[w["key"]][w["line"]].append(w)
    out = []
    for key in order:
        lines = [paras[key][ln] for ln in sorted(paras[key])]
        text = ""
        for ln in lines:
            chunk = " ".join(w["text"] for w in ln)
            if text.endswith("-") and chunk[:1].islower():
                text = text[:-1] + chunk          # re-join hyphenated line breaks
            else:
                text = (text + " " + chunk).strip()
        heights = [statistics.median(w["h"] for w in ln) for ln in lines]
        line_texts = [" ".join(w["text"] for w in ln) for ln in lines]
        line_words = [[w["id"] for w in ln] for ln in lines]
        out.append({"text": text, "lines": len(lines), "line_texts": line_texts,
                    "line_words": line_words, "h": statistics.median(heights)})
    return out


def _text_layer_lines(pdf: str) -> list[str]:
    """Lines from the PDF's own (often hidden OCR) text layer, first pages only, for the TOC."""
    try:
        import pymupdf
        with pymupdf.open(pdf) as d:
            text = "\n".join(d[i].get_text("text") for i in range(min(6, len(d))))
        return [l for l in text.splitlines() if l.strip()]
    except Exception:
        return []


DPI = 300


def analyze_scan(pdf: str, lang: str, title: str, use_heading_title: bool = True,
                 font_dir: str | None = None, cancel_file: str = "") -> dict:
    """OCR every page once; return words (with positions), structure items and the title.

    If font_dir is given, Tesseract also writes a text-only PDF for page 1 there, whose
    invisible embedded font is reused for the tagged text layer."""
    work = tempfile.mkdtemp(prefix="ocrdoc_")
    words_by_id: dict[int, dict] = {}
    pages = []
    font_pdf = None
    try:
        from .common import run_cmd
        run_cmd(["pdftoppm", "-r", str(DPI), "-gray", "-png", pdf, os.path.join(work, "p")],
                       check=True, capture_output=True, timeout=3600)
        pngs = sorted(glob.glob(os.path.join(work, "p-*.png")))
        next_id = 0
        page_info = []
        for n, png in enumerate(pngs, 1):
            if cancel_file and os.path.exists(cancel_file):
                from .common import Cancelled
                raise Cancelled()
            words, turn, conf = ocr_upright(png, lang)
            page_info.append({"page": n, "turn": turn, "conf": round(conf, 1), "words": len(words),
                              **ocr_review(png, turn, words)})
            for w in words:
                w["id"] = next_id
                w["page"] = n
                words_by_id[next_id] = w
                next_id += 1
            paras = _paragraphs(words)
            for p in paras:
                p["page"] = n
            pages.append(paras)
        if font_dir and pngs:
            base = os.path.join(font_dir, "glyphless")
            subprocess.run(["tesseract", _rotated_copy(pngs[0], page_info[0]["turn"]), base, "-l", lang, "--dpi", str(DPI),
                            "-c", "textonly_pdf=1", "pdf"], capture_output=True, timeout=600)
            font_pdf = base + ".pdf" if os.path.exists(base + ".pdf") else None
    finally:
        subprocess.run(["rm", "-rf", work])

    all_paras = [p for pg in pages for p in pg if len(p["text"]) >= 2]
    body_h = statistics.median([p["h"] for p in all_paras if p["lines"] > 1] or
                               [p["h"] for p in all_paras] or [30])

    def conf_of(p):
        ws = [words_by_id[w] for l in p["line_words"] for w in l]
        return sum(w["conf"] for w in ws) / len(ws) if ws else 0

    def wordy(text):
        letters = sum(c.isalpha() for c in text)
        visible = sum(not c.isspace() for c in text) or 1
        return letters >= 4 and letters / visible >= 0.7

    # title: the largest short, clean paragraph on page 1 (a chapter title beats a leftover "Note")
    title_para = None
    cands = [p for p in all_paras if p["page"] == 1 and len(p["text"].split()) <= 25
             and p["h"] >= body_h * 1.1 and not p["text"].endswith((".", ","))
             and wordy(p["text"]) and conf_of(p) >= 75
             and (p["lines"] <= 3 or p["h"] >= body_h * 1.8)]
    if cands:
        title_para = max(cands, key=lambda p: p["h"])
    rest = [p for p in all_paras if p is not title_para]
    for p in rest:
        p["big"] = p["lines"] <= 2 and p["h"] >= body_h * 1.3
        p["lines"] = p["line_texts"]
    items, info = build_structure(rest, len(pages), layer_lines=_text_layer_lines(pdf))
    # a heading must read like words: no OCR debris, table cells, symbols or "(Continued)"
    for it in items:
        if it["kind"] != "heading" or not it.get("words"):
            continue
        ws = [words_by_id[w] for w in it["words"]]
        conf = sum(w["conf"] for w in ws) / len(ws)
        t = it["text"].strip()
        if conf < 80 or not wordy(t) or t.startswith("(") or t.lower() in {"totals", "total"}:
            it["kind"] = "para"
    if title_para is not None:
        items.insert(0, {"kind": "heading", "text": title_para["text"], "level": 1,
                         "words": [w for l in title_para["line_words"] for w in l]})
        if use_heading_title:
            title = title_para["text"][:200]
    elif use_heading_title:
        first = next((i["text"] for i in items if i["kind"] == "heading"), None)
        if first:
            title = first[:200]
    title = " ".join(title.split())
    return {"pages": len(pages), "items": items, "words": words_by_id, "title": title,
            "structure": info["source"], "toc_entries": info["toc_entries"], "font_pdf": font_pdf,
            "headings": sum(1 for i in items if i["kind"] == "heading"), "page_info": page_info}


def write_text_docx(scan: dict, out_docx: str) -> dict:
    doc = Document()
    doc.core_properties.title = scan["title"]
    n_head = n_para = 0
    from .structure import normalize_heading_levels
    normalize_heading_levels(scan["items"])
    for it in scan["items"]:
        if it["kind"] == "heading":
            doc.add_heading(it["text"], level=min(max(it["level"], 1), 6))
            n_head += 1
        else:
            para = doc.add_paragraph(clean_leaders(it["text"]))
            para.paragraph_format.space_after = Pt(8)
            n_para += 1
    from .contrast import fix_docx_contrast
    fix_docx_contrast(doc)
    doc.save(out_docx)
    return {"headings": n_head, "paragraphs": n_para}


def scanned_pdf_to_docx(pdf: str, out_docx: str, lang: str, title: str,
                        use_heading_title: bool = True) -> dict:
    scan = analyze_scan(pdf, lang, title, use_heading_title)
    counts = write_text_docx(scan, out_docx)
    return {"pages": scan["pages"], "title": scan["title"], "structure": scan["structure"],
            "toc_entries": scan["toc_entries"], **counts}
